"""Render aproximado do .pptx a partir das posicoes reais do XML.

NAO e o que o PowerPoint mostra. Serve pra achar: texto estourando a caixa,
shape fora do slide, cor errada, elemento faltando.
"""
import sys, os
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont

SRC = sys.argv[1]
OUT = sys.argv[2]
os.makedirs(OUT, exist_ok=True)

prs = Presentation(SRC)
SW, SH = prs.slide_width, prs.slide_height
PX = 1280.0 / SW  # escala EMU -> px

def font_path(bold=False, serif=False, mono=False):
    base = r"C:\Windows\Fonts"
    if mono:
        return os.path.join(base, "consola.ttf")
    if serif:
        return os.path.join(base, "georgiab.ttf" if bold else "georgia.ttf")
    return os.path.join(base, "arialbd.ttf" if bold else "arial.ttf")

def load(sz, bold=False, serif=False, mono=False):
    try:
        return ImageFont.truetype(font_path(bold, serif, mono), max(8, int(sz)))
    except Exception:
        return ImageFont.load_default()

def rgb(c, default=(120, 120, 120)):
    try:
        if c and c.type is not None and c.rgb is not None:
            return tuple(bytes.fromhex(str(c.rgb)))
    except Exception:
        pass
    return default

problems = []

for i, slide in enumerate(prs.slides, 1):
    W, H = int(SW * PX), int(SH * PX)
    try:
        base = tuple(bytes.fromhex(str(slide.background.fill.fore_color.rgb)))
    except Exception:
        base = (255, 255, 255)
    img = Image.new("RGB", (W, H), base)
    d = ImageDraw.Draw(img)

    for sh in slide.shapes:
        try:
            x, y = int(sh.left * PX), int(sh.top * PX)
            w, h = int(sh.width * PX), int(sh.height * PX)
        except Exception:
            continue

        # fora dos limites?
        if x < -2 or y < -2 or x + w > W + 2 or y + h > H + 2:
            problems.append(f"slide {i}: '{sh.shape_type}' fora dos limites "
                            f"({x},{y},{x+w},{y+h}) vs (0,0,{W},{H})")

        # fundo
        fill = None
        try:
            if sh.fill.type is not None and sh.fill.type == 1:
                fill = rgb(sh.fill.fore_color, None)
        except Exception:
            pass
        if fill:
            d.rectangle([x, y, x + w, y + h], fill=fill)
        elif sh.shape_type is not None and "PICTURE" in str(sh.shape_type):
            d.rectangle([x, y, x + w, y + h], outline=(200, 0, 200), width=2)
            d.text((x + 4, y + 4), "[img]", fill=(200, 0, 200), font=load(12))

        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue

        cy = y + 4
        for p in sh.text_frame.paragraphs:
            txt = "".join(r.text for r in p.runs)
            if not txt.strip():
                cy += 10
                continue
            r0 = p.runs[0]
            sz = (r0.font.size.pt if r0.font.size else
                  (p.font.size.pt if p.font.size else 18))
            name = (r0.font.name or "").lower()
            f = load(sz * 12700 * PX, bold=bool(r0.font.bold),
                     serif="georgia" in name, mono="courier" in name)
            col = rgb(r0.font.color, (30, 30, 30))

            # quebra de linha manual na largura da caixa
            words, line = txt.split(), ""
            lines = []
            for wd in words:
                t = (line + " " + wd).strip()
                if d.textlength(t, font=f) > max(20, w - 8) and line:
                    lines.append(line)
                    line = wd
                else:
                    line = t
            if line:
                lines.append(line)
            for ln in lines:
                d.text((x + 4, cy), ln, fill=col, font=f)
                cy += int(sz * 12700 * PX * 1.28)

        if cy > y + h + 6:
            problems.append(f"slide {i}: texto passa da caixa em "
                            f"{int(cy - (y+h))}px -> "
                            f"\"{sh.text_frame.text[:45].strip()}...\"")
        if cy > H + 2:
            problems.append(f"slide {i}: texto sai do slide (y={cy} > {H})")

    img.save(os.path.join(OUT, f"slide-{i}.png"))

print(f"slides: {len(prs.slides)}  tamanho: {SW}x{SH} EMU  "
      f"({SW/914400:.2f}x{SH/914400:.2f} pol, "
      f"{'16:9' if abs(SW/SH - 16/9) < 0.01 else 'NAO 16:9'})")
print("-" * 60)
if problems:
    for p in problems:
        print("!", p)
else:
    print("nenhum problema geometrico detectado")
