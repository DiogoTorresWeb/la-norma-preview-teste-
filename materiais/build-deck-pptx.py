"""Gera a versao .pptx editavel do deck de 5 slides da La Norma.

Espelha materiais/deck.html. Como o PowerPoint so usa fontes instaladas na
maquina, Fraunces/Archivo/IBM Plex Mono sao substituidas por Georgia/Arial/
Courier New -- mesma logica de fallback usada na plantilla de e-mail.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------- tokens (iguais aos do site) ----------
ESPRESSO_DEEP = RGBColor(0x17, 0x0F, 0x0A)
ESPRESSO      = RGBColor(0x24, 0x1C, 0x15)
CREAM         = RGBColor(0xF3, 0xEC, 0xDD)
SAND          = RGBColor(0xE7, 0xD9, 0xBE)
BRASS         = RGBColor(0xB8, 0x84, 0x3C)
BRASS_LIGHT   = RGBColor(0xE4, 0xC4, 0x8C)
MUTED         = RGBColor(0xB7, 0xA4, 0x89)
MUTED_DARK    = RGBColor(0x6B, 0x5B, 0x45)
BAD_DARK      = RGBColor(0xE0, 0x80, 0x74)
BAD_LIGHT     = RGBColor(0xA8, 0x39, 0x2F)
GOOD_LIGHT    = RGBColor(0x3F, 0x6B, 0x46)

F_DISPLAY = "Georgia"
F_BODY    = "Arial"
F_MONO    = "Courier New"

W, H = 13.333, 7.5
MARGIN = 0.9
CONTENT_W = W - MARGIN * 2

LOGO_CREAM = "materiais/assets/logo-lanorma-cream.png"
LOGO_DARK  = "assets/real/logo-lanorma.png"
COVER_IMG  = "assets/real/hero-display.jpg"

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def text(slide, s, x, y, w, h, size, color, font=F_BODY, bold=False,
         spacing=None, caps=False, align=PP_ALIGN.LEFT, line=None, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    if line:
        p.line_spacing = line
    run = p.add_run()
    run.text = s.upper() if caps else s
    f = run.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    if spacing is not None:
        # letter-spacing nao tem API publica: escreve o atributo spc no XML
        run.font._rPr.set('spc', str(int(spacing * 100)))
    return box


def head(slide, num, dark=True, show_logo=True):
    if show_logo:
        logo = LOGO_CREAM if dark else LOGO_DARK
        slide.shapes.add_picture(logo, Inches(MARGIN), Inches(0.46), height=Inches(0.23))
    text(slide, num, W - MARGIN - 2, 0.47, 2, 0.3, 10, MUTED if dark else MUTED_DARK,
         font=F_MONO, align=PP_ALIGN.RIGHT, spacing=0.6)
    rect(slide, MARGIN, 0.92, CONTENT_W, 0.012,
         RGBColor(0x3A, 0x2E, 0x24) if dark else RGBColor(0xD6, 0xC9, 0xB2))


def foot(slide, left_plain, left_bold, right, dark=True):
    rect(slide, MARGIN, 6.62, CONTENT_W, 0.012,
         RGBColor(0x3A, 0x2E, 0x24) if dark else RGBColor(0xD6, 0xC9, 0xB2))
    box = slide.shapes.add_textbox(Inches(MARGIN), Inches(6.78), Inches(CONTENT_W * 0.62), Inches(0.45))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = left_plain + " "
    r1.font.name = F_BODY; r1.font.size = Pt(10.5); r1.font.color.rgb = MUTED if dark else MUTED_DARK
    r2 = p.add_run(); r2.text = left_bold
    r2.font.name = F_BODY; r2.font.size = Pt(10.5); r2.font.bold = True
    r2.font.color.rgb = CREAM if dark else ESPRESSO
    text(slide, right, W - MARGIN - 5, 6.8, 5, 0.4, 9.5, MUTED if dark else MUTED_DARK,
         font=F_MONO, align=PP_ALIGN.RIGHT, spacing=0.3)


def eyebrow(slide, s, y, dark=True):
    text(slide, s, MARGIN, y, 6, 0.28, 10.5, BRASS if dark else MUTED_DARK,
         bold=True, caps=True, spacing=2.0)


def title(slide, s, y, size=34, dark=True, w=None):
    text(slide, s, MARGIN, y, w or CONTENT_W * 0.72, 1.9, size,
         CREAM if dark else ESPRESSO, font=F_DISPLAY, bold=True, line=1.06)


def lede(slide, s, y, dark=True, w=None):
    text(slide, s, MARGIN, y, w or CONTENT_W * 0.62, 0.9, 13,
         MUTED if dark else MUTED_DARK, line=1.5)


def point(slide, x, y, w, step, h3, body, dark=True):
    """Bloco editorial: regua brass no topo, sem caixa (padrao do site)."""
    rect(slide, x, y, w, 0.028, BRASS)
    text(slide, step, x, y + 0.22, w, 0.26, 10.5, BRASS, font=F_MONO, spacing=1.0)
    text(slide, h3, x, y + 0.56, w, 0.6, 16, CREAM if dark else ESPRESSO,
         font=F_DISPLAY, bold=True, line=1.18)
    text(slide, body, x, y + 1.18, w, 1.5, 11.5, MUTED if dark else MUTED_DARK, line=1.5)


# =====================================================================
# 01 · PORTADA
# =====================================================================
s1 = prs.slides.add_slide(BLANK)
bg(s1, ESPRESSO_DEEP)
# foto de fundo na metade direita, recortada como object-fit:cover
PIC_X = 5.1
PIC_W = W - PIC_X                      # 8.233 in
pic = s1.shapes.add_picture(COVER_IMG, Inches(PIC_X), Inches(0),
                            width=Inches(PIC_W), height=Inches(H))
# a foto e 3:2; a caixa e mais alta, entao sobra largura -> corta nas laterais,
# puxando o enquadramento para a direita (onde esta o display aceso)
_visible = (PIC_W / H) / (2400 / 1600)
_cut = 1 - _visible
pic.crop_left = _cut * 0.75
pic.crop_right = _cut * 0.25
veil = rect(s1, 5.1, 0, W - 5.1, H, ESPRESSO_DEEP)
veil.fill.fore_color.rgb = ESPRESSO_DEEP
veil.fill.transparency = 0.45          # ignorado por algumas versoes; ver alpha abaixo
# alpha real no XML (transparency nao e exposto de forma fiavel pela API)
from pptx.oxml.ns import qn
solid = veil.fill._xPr.find(qn('a:solidFill'))
clr = solid.find(qn('a:srgbClr'))
alpha = clr.makeelement(qn('a:alpha'), {'val': '55000'})
clr.append(alpha)
# faixa solida a esquerda para o texto respirar
rect(s1, 0, 0, 5.4, H, ESPRESSO_DEEP)

head(s1, "01 / 05", dark=True, show_logo=False)
s1.shapes.add_picture(LOGO_CREAM, Inches(MARGIN), Inches(2.15), height=Inches(0.42))
eyebrow(s1, "Propuesta interna · Septiembre 2026", 2.95)
title(s1, "Puedo acelerar la web para que venda al ritmo del taller.", 3.35, size=30, w=4.2)
text(s1, "Una propuesta directa para resolver el principal cuello de botella "
         "digital de La Norma hoy — hecha por quien monta las máquinas desde dentro.",
     MARGIN, 5.35, 4.1, 1.0, 11.5, MUTED, line=1.5)
foot(s1, "", "Diogo Torres", "Palma de Gandia, Valencia", dark=True)
text(s1, "Montaje en taller & desarrollo web", MARGIN, 7.0, 4, 0.3, 10.5, MUTED)

# =====================================================================
# 02 · DIAGNÓSTICO
# =====================================================================
s2 = prs.slides.add_slide(BLANK)
bg(s2, CREAM)
head(s2, "02 / 05", dark=False)
eyebrow(s2, "Diagnóstico", 1.45, dark=False)
title(s2, "La máquina es excelente en la barra. La web frena en el móvil.",
      1.85, size=30, dark=False, w=8.6)
lede(s2, "Comprobado con las herramientas oficiales de auditoría. La mayoría de "
         "hosteleros y distribuidores entra desde el teléfono.", 3.3, dark=False, w=7.6)

DATA = [
    ("35", "/100", BAD_LIGHT, "PageSpeed en móvil",
     "64 sobre 100 en ordenador. El catálogo tarda demasiado en mostrar fotos y vídeos sin comprimir."),
    ("0", "", ESPRESSO, "Formularios de contacto",
     "Quien quiere cotizar tiene que buscar el teléfono o abrir su propio correo. Cada clic extra enfría la venta."),
    ("4,1", "/5", BAD_LIGHT, "Google Business · 8 reseñas",
     "La más reciente es de 1 estrella: pidió información y no tuvo respuesta clara ni rápida."),
]
col_w = CONTENT_W / 3
rect(s2, MARGIN, 4.22, CONTENT_W, 0.012, RGBColor(0xD6, 0xC9, 0xB2))
for i, (val, unit, col, label, note) in enumerate(DATA):
    x = MARGIN + i * col_w
    if i:
        rect(s2, x - 0.01, 4.22, 0.012, 1.95, RGBColor(0xD6, 0xC9, 0xB2))
    box = s2.shapes.add_textbox(Inches(x + (0.28 if i else 0)), Inches(4.5),
                                Inches(col_w - 0.5), Inches(0.6))
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = val
    r.font.name = F_MONO; r.font.size = Pt(32); r.font.color.rgb = col
    if unit:
        r2 = p.add_run(); r2.text = unit
        r2.font.name = F_MONO; r2.font.size = Pt(16); r2.font.color.rgb = MUTED_DARK
    text(s2, label, x + (0.28 if i else 0), 5.15, col_w - 0.5, 0.3, 9.5, MUTED_DARK,
         bold=True, caps=True, spacing=1.4)
    text(s2, note, x + (0.28 if i else 0), 5.5, col_w - 0.5, 1.1, 10.5, MUTED_DARK, line=1.45)
rect(s2, MARGIN, 6.17, CONTENT_W, 0.012, RGBColor(0xD6, 0xC9, 0xB2))
foot(s2, "La Norma no tiene un problema de producto.", "Tiene un problema de fricción en la entrada.",
     "Fuente: PageSpeed Insights y Google Maps · sept. 2026", dark=False)

# =====================================================================
# 03 · PRUEBA TÉCNICA
# =====================================================================
s3 = prs.slides.add_slide(BLANK)
bg(s3, ESPRESSO_DEEP)
head(s3, "03 / 05")
eyebrow(s3, "Prueba técnica", 1.45)
title(s3, "De 35 a 99 sobre 100. Sé exactamente cómo se hace.", 1.85, size=30, w=8.4)
lede(s3, "No es teoría de agencia: son dos sitios que ya existen y se pueden abrir "
         "en directo, ahora mismo, en este portátil.", 3.25, w=7.6)

half = CONTENT_W / 2 - 0.4
COMPARE = [
    (MARGIN, "HOY", "lanorma.es — WordPress con sobrecarga",
     [("Rendimiento móvil", "35 / 100", BAD_DARK),
      ("Rendimiento ordenador", "64 / 100", BAD_DARK),
      ("Carga del catálogo", "lenta", BAD_DARK)]),
    (MARGIN + half + 0.8, "HECHO POR MÍ", "oliveirafotos.es y el prototipo Ln 200",
     [("Rendimiento móvil", "99 / 100", BRASS_LIGHT),
      ("Rendimiento ordenador", "100 / 100", BRASS_LIGHT),
      ("Carga en teléfono", "< 0,6 s", BRASS_LIGHT)]),
]
for x, step, h3, rows in COMPARE:
    rect(s3, x, 4.08, half, 0.028, BRASS)
    text(s3, step, x, 4.3, half, 0.26, 10.5, BRASS, font=F_MONO, spacing=1.0)
    text(s3, h3, x, 4.62, half, 0.5, 15, CREAM, font=F_DISPLAY, bold=True, line=1.2)
    for j, (lab, val, col) in enumerate(rows):
        y = 5.25 + j * 0.38
        text(s3, lab, x, y, half * 0.62, 0.3, 11, MUTED)
        text(s3, val, x + half * 0.5, y, half * 0.5, 0.3, 11, col,
             font=F_MONO, align=PP_ALIGN.RIGHT)
        rect(s3, x, y + 0.3, half, 0.008, RGBColor(0x3A, 0x2E, 0x24))
text(s3, "Una agencia externa cobraría entre 2.500 € y 4.000 €, y tardaría dos meses "
         "solo en entender qué máquina vendemos y qué componentes lleva.",
     MARGIN, 6.62 - 0.72, CONTENT_W * 0.74, 0.6, 13, CREAM,
     font=F_DISPLAY, italic=True, line=1.4)
foot(s3, "Prueba en vivo:", "podemos abrir la comparativa en el navegador cuando quieras.",
     "Google Lighthouse · Core Web Vitals")

# =====================================================================
# 04 · PROPUESTA
# =====================================================================
s4 = prs.slides.add_slide(BLANK)
bg(s4, SAND)
head(s4, "04 / 05", dark=False)
eyebrow(s4, "Propuesta", 1.45, dark=False)
title(s4, "Dejar la web impecable y lista para vender.", 1.85, size=30, dark=False, w=8.0)
lede(s4, "Un trabajo quirúrgico y directo, empezando por lo que genera impacto "
         "inmediato en la venta.", 3.15, dark=False, w=7.2)

POINTS4 = [
    ("01", "Aligerar el catálogo",
     "Que el cliente vea la gama Ln1 —Compacta, 2 y 3 grupos— y el lanzamiento Ln 200 "
     "sin esperar un solo segundo en el móvil."),
    ("02", "Canal directo al taller",
     "«Dinos cuántos cafés sirves al día y te decimos qué máquina necesitas.» "
     "Contacto en un clic, sin buscar el teléfono."),
    ("03", "Ya estoy dentro",
     "Monto las máquinas con mis propias manos. Sé qué es el grupo erogador, el PID "
     "y la caldera: cero tiempo de adaptación."),
]
pw = (CONTENT_W - 0.9) / 3
for i, (step, h3, body) in enumerate(POINTS4):
    point(s4, MARGIN + i * (pw + 0.45), 3.95, pw, step, h3, body, dark=False)
text(s4, "Si una sola máquina se vende porque un distribuidor navegó cómodo en el móvil "
         "y dejó sus datos, la web queda amortizada varias veces.",
     MARGIN, 6.05, CONTENT_W * 0.74, 0.5, 13, ESPRESSO, font=F_DISPLAY, italic=True, line=1.4)
foot(s4, "Ejecución rápida:", "prototipo listo para integrar en días, sin interrumpir el taller.",
     "Cero riesgo · retorno inmediato", dark=False)

# =====================================================================
# 05 · SIGUIENTE PASO
# =====================================================================
s5 = prs.slides.add_slide(BLANK)
bg(s5, ESPRESSO_DEEP)
head(s5, "05 / 05")
eyebrow(s5, "Siguiente paso", 1.45)
title(s5, "Primero la web. Después, lo que La Norma vaya necesitando.", 1.85, size=28, w=8.4)

POINTS5 = [
    ("PASO 1 · AHORA", "La web y el catálogo",
     "Resolver la velocidad en móvil, publicar la ficha de la Ln 200 y abrir el canal "
     "de contacto para cotizaciones directas."),
    ("PASO 2", "Redes con cara y ojos",
     "Alimentar LinkedIn e Instagram con el taller real: la ingeniería y el montaje "
     "que la competencia no puede enseñar."),
    ("PASO 3", "Seguimiento de ventas",
     "Que ningún presupuesto enviado se quede sin respuesta, y que los avisos no "
     "desborden la bandeja de entrada."),
]
for i, (step, h3, body) in enumerate(POINTS5):
    point(s5, MARGIN + i * (pw + 0.45), 3.3, pw, step, h3, body, dark=True)

rect(s5, MARGIN, 5.5, CONTENT_W, 0.028, BRASS)
text(s5, "Hablamos 10 minutos cuando te venga bien.", MARGIN, 5.75, 7.4, 0.5, 20,
     CREAM, font=F_DISPLAY, bold=True)
text(s5, "Te enseño el prototipo en el móvil o en el portátil, en un descanso del taller.",
     MARGIN, 6.2, 7.4, 0.35, 11.5, MUTED)
chip = rect(s5, W - MARGIN - 3.5, 5.78, 3.5, 0.52, ESPRESSO_DEEP)
chip.line.color.rgb = BRASS
chip.line.width = Pt(1)
text(s5, "Diogo Torres · Taller de Valencia", W - MARGIN - 3.5, 5.95, 3.5, 0.3, 10.5,
     BRASS_LIGHT, font=F_MONO, align=PP_ALIGN.CENTER, spacing=0.4)
foot(s5, "Propuesta interna:", "hecha para sumar a La Norma.",
     "La Norma Coffee Machine Manufacturer S.L.")

OUT = "materiais/La-Norma-Propuesta-2026.pptx"
prs.save(OUT)
import os
print(f"{OUT}  {os.path.getsize(OUT)} bytes  {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
